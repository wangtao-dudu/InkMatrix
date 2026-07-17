# -*- coding: utf-8 -*-
"""
product_page.py  (UI层 - 产品管理 / 配色模板库)
======================================================
这一页只管一件事：**这个产品是怎么配色的。**

    产品名 / 客户 / 承印物 / 有几个色位 / 每个色位的目标色

── 为什么这里没有钱的事了 ──────────────────────────────────
以前这一页混了一大堆财务字段：数量、单价、优惠、实际应付、已付、结余、
欠款周期……问题在于，这些数字在【多色印刷工单】页也要录一遍，在【进销存】
里还有一份。同一笔生意的钱，在三个地方各存一份，靠人去保证一致——迟早对
不上，而且对不上的时候你根本不知道该信哪一个。

现在钱的事统一归【财务管理】：它直接挂在进销存的进货记录上，那才是钱的
源头（客户送来多少货 → 按单价算多少钱 → 收了多少 → 还欠多少）。

产品模板则只管配色。同一个配色模板可以复用在很多笔生意上——客户 A 的
100ml 白瓶三色套印，这个月做一批，下个月再做一批，配色是同一套，钱是两笔账。
把它们绑死在一起本来就是错的。

── 模板怎么来的 ──────────────────────────────────────────
在【多色印刷工单】页点【全部色位智能寻优】，只要填了产品名称，寻优成功
后会**自动**存成模板，不用手动点保存（以前忘了点就白算一遍）。
"""
import datetime
from typing import List, Dict, Optional

from PyQt6.QtCore import Qt, pyqtSignal
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit,
    QTableWidget, QTableWidgetItem, QHeaderView, QMessageBox,
    QFileDialog, QAbstractItemView, QDialog, QTextEdit, QDialogButtonBox,
)

from core_layer.color_convert import lab_to_srgb_approx
from ui_layer.widgets import Card, PrimaryButton, GhostButton, DangerButton, StatCard


class ProductPage(QWidget):

    load_requested = pyqtSignal(dict)   # 请求把某个产品模板载入到工单页
    status_message = pyqtSignal(str)

    def __init__(self, product_repo, bottle_repo, parent=None):
        super().__init__(parent)
        self.product_repo = product_repo
        self.bottle_repo = bottle_repo
        self._build_ui()
        self.refresh_table()

    # ------------------------------------------------------------------

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 20, 24, 20)
        root.setSpacing(14)

        title = QLabel("产品管理 · 配色模板库")
        title.setObjectName("PageTitle")
        subtitle = QLabel(
            "已保存的产品配色方案：承印物 + 各个色位的目标色。"
            "在【多色印刷工单】寻优成功后会自动存进来，下次同样的活直接载入，不用重新算"
        )
        subtitle.setObjectName("PageSubtitle")
        root.addWidget(title)
        root.addWidget(subtitle)

        stats = QHBoxLayout()
        stats.setSpacing(12)
        self.stat_total = StatCard("0", "配色模板总数", accent_hex="#0d9488")
        self.stat_customers = StatCard("0", "涉及客户数", accent_hex="#2563eb")
        self.stat_stations = StatCard("0", "色位总数", accent_hex="#7c3aed")
        for s in (self.stat_total, self.stat_customers, self.stat_stations):
            stats.addWidget(s)
        root.addLayout(stats)

        money_note = QLabel(
            "💰 数量 / 单价 / 优惠 / 收款等财务信息，请到【财务管理】页查看——"
            "那边的数据直接来自【进销存管理】的进货记录，是钱的唯一源头。"
            "配色模板可以在很多笔生意上重复用，跟某一笔具体的账不是一回事，所以分开管。"
        )
        money_note.setWordWrap(True)
        money_note.setStyleSheet(
            "color:#0369a1; background:#f0f9ff; border:1px solid #bae6fd; "
            "border-radius:8px; padding:9px 12px; font-size:12px;"
        )
        root.addWidget(money_note)

        card = Card("配色模板列表")
        lay = card.body_layout

        bar = QHBoxLayout()
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("搜索产品名 / 客户 / 电话…")
        self.search_edit.textChanged.connect(self.refresh_table)
        bar.addWidget(self.search_edit, 2)
        bar.addStretch(1)

        detail_btn = GhostButton("🎨 查看色位明细")
        detail_btn.clicked.connect(self._on_view_stations)
        load_btn = PrimaryButton("📂 载入到工单")
        load_btn.clicked.connect(self._on_load)
        export_btn = GhostButton("📊 导出Excel")
        export_btn.clicked.connect(self._on_export_excel)
        del_btn = DangerButton("－ 删除模板")
        del_btn.setMinimumHeight(34)
        del_btn.clicked.connect(self._on_delete)
        for b in (detail_btn, export_btn, load_btn, del_btn):
            bar.addWidget(b)
        lay.addLayout(bar)

        self.table = QTableWidget(0, 8)
        self.table.setHorizontalHeaderLabels([
            "产品编号", "产品名称", "客户", "联系电话",
            "承印物", "色位数", "色位预览", "更新时间",
        ])
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        h.resizeSection(6, 160)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setDefaultSectionSize(36)
        self.table.doubleClicked.connect(self._on_load)
        lay.addWidget(self.table, 1)

        hint = QLabel("💡 双击任意一行可以直接把这个配色模板载入到【多色印刷工单】。")
        hint.setStyleSheet("color:#64748b; font-size:11px;")
        lay.addWidget(hint)

        root.addWidget(card, 1)

    # ------------------------------------------------------------------

    def _bottle_name(self, code: str) -> str:
        if not code:
            return "-"
        rec = self.bottle_repo.get_by_code(code) if self.bottle_repo else None
        return rec["name"] if rec else code

    def _filtered_records(self) -> List[Dict]:
        records = self.product_repo.load_all()
        kw = self.search_edit.text().strip().lower()
        if not kw:
            return records
        out = []
        for r in records:
            blob = " ".join([
                r.get("name", "") or "", r.get("customer", "") or "",
                r.get("customer_phone", "") or "", r.get("code", "") or "",
            ]).lower()
            if kw in blob:
                out.append(r)
        return out

    def refresh_table(self):
        records = self._filtered_records()
        self.table.setRowCount(len(records))

        customers = set()
        station_total = 0

        for i, r in enumerate(records):
            stations = r.get("stations", []) or []
            station_total += len(stations)
            if r.get("customer"):
                customers.add(r["customer"])

            cells = [
                r.get("code", ""),
                r.get("name", ""),
                r.get("customer", "") or "-",
                r.get("customer_phone", "") or "-",
                self._bottle_name(r.get("bottle_code", "")),
                str(len(stations)),
                "",   # 色位预览：下面单独塞色块
                r.get("updated_at", "") or "-",
            ]
            for col, text in enumerate(cells):
                item = QTableWidgetItem(text)
                if col == 0:
                    item.setData(Qt.ItemDataRole.UserRole, r.get("code", ""))
                if col != 1:
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(i, col, item)

            self.table.setCellWidget(i, 6, self._make_station_preview(stations))

        self.stat_total.set_value(str(len(records)))
        self.stat_customers.set_value(str(len(customers)))
        self.stat_stations.set_value(str(station_total))

    def _make_station_preview(self, stations: List[Dict]) -> QWidget:
        """把这个产品的各个色位画成一排小色块——一眼就能认出是哪个活。"""
        w = QWidget()
        lay = QHBoxLayout(w)
        lay.setContentsMargins(6, 4, 6, 4)
        lay.setSpacing(3)
        for s in stations[:8]:
            lab = s.get("target_lab") or [50, 0, 0]
            try:
                r, g, b = lab_to_srgb_approx(*lab)
            except Exception:
                r, g, b = 200, 200, 200
            chip = QLabel()
            chip.setFixedSize(18, 18)
            chip.setStyleSheet(
                f"background: rgb({r},{g},{b}); border:1px solid #cbd5e1; border-radius:3px;"
            )
            chip.setToolTip(f"{s.get('name','')}　Lab({lab[0]:.0f}, {lab[1]:.0f}, {lab[2]:.0f})")
            lay.addWidget(chip)
        if len(stations) > 8:
            more = QLabel(f"+{len(stations) - 8}")
            more.setStyleSheet("color:#64748b; font-size:10px;")
            lay.addWidget(more)
        lay.addStretch(1)
        return w

    # ------------------------------------------------------------------

    def _selected_record(self) -> Optional[Dict]:
        row = self.table.currentRow()
        if row < 0:
            QMessageBox.information(self, "提示", "请先在表格中点击选中一个产品模板。")
            return None
        code = self.table.item(row, 0).data(Qt.ItemDataRole.UserRole)
        rec = self.product_repo.get_by_code(code) if hasattr(self.product_repo, "get_by_code") else None
        if rec is None:
            for r in self.product_repo.load_all():
                if r.get("code") == code:
                    rec = r
                    break
        return rec

    def _on_view_stations(self):
        rec = self._selected_record()
        if not rec:
            return
        stations = rec.get("stations", []) or []
        if not stations:
            QMessageBox.information(self, "提示", "这个模板里没有色位。")
            return

        lines = [
            f"产品：{rec.get('name','')}",
            f"客户：{rec.get('customer','') or '-'}",
            f"承印物：{self._bottle_name(rec.get('bottle_code',''))}",
            "", "色位明细：",
        ]
        for i, s in enumerate(stations, 1):
            lab = s.get("target_lab") or [0, 0, 0]
            lines.append(
                f"  {i}. {s.get('name','')}　"
                f"Lab({lab[0]:.2f}, {lab[1]:.2f}, {lab[2]:.2f})　"
                f"配墨 {s.get('total_grams', 0):g} g"
            )
        lines += ["", "💡 载入到工单后，需要重新点【全部色位智能寻优】——",
                  "   因为油墨库存和可用油墨可能已经变了，旧配方不一定还适用。"]

        dlg = QDialog(self)
        dlg.setWindowTitle("色位明细")
        dlg.setMinimumSize(480, 380)
        v = QVBoxLayout(dlg)
        te = QTextEdit()
        te.setReadOnly(True)
        te.setPlainText("\n".join(lines))
        v.addWidget(te)
        b = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        b.rejected.connect(dlg.reject)
        b.accepted.connect(dlg.accept)
        v.addWidget(b)
        dlg.exec()

    def _on_load(self):
        rec = self._selected_record()
        if not rec:
            return
        self.load_requested.emit(rec)

    def _on_delete(self):
        rec = self._selected_record()
        if not rec:
            return
        confirm = QMessageBox.question(
            self, "确认删除",
            f"确定要删除配色模板【{rec.get('name','')}】吗？\n\n"
            f"只是删掉这套配色方案，不影响任何进销存记录和财务账目。"
        )
        if confirm != QMessageBox.StandardButton.Yes:
            return
        self.product_repo.delete(rec.get("code", ""))
        self.refresh_table()
        self.status_message.emit(f"已删除配色模板【{rec.get('name','')}】")

    def _on_export_excel(self):
        records = self._filtered_records()
        if not records:
            QMessageBox.information(self, "提示", "当前没有可导出的产品模板。")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "导出配色模板",
            f"配色模板_{datetime.date.today().isoformat()}.xlsx",
            "Excel 文件 (*.xlsx)"
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            QMessageBox.warning(self, "缺少依赖", "导出Excel需要 openpyxl，请先安装：pip install openpyxl")
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "配色模板"
        headers = ["产品编号", "产品名称", "客户", "联系电话", "承印物",
                   "色位名称", "目标L*", "目标a*", "目标b*", "配墨量(g)", "更新时间"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0D9488")
            c.alignment = Alignment(horizontal="center")

        # 一个产品有几个色位就展开成几行——这样导出的表可以直接拿去打印给现场
        for r in records:
            stations = r.get("stations", []) or []
            if not stations:
                ws.append([r.get("code",""), r.get("name",""), r.get("customer",""),
                           r.get("customer_phone",""), self._bottle_name(r.get("bottle_code","")),
                           "(无色位)", "", "", "", "", r.get("updated_at","")])
                continue
            for s in stations:
                lab = s.get("target_lab") or [0, 0, 0]
                ws.append([
                    r.get("code",""), r.get("name",""), r.get("customer",""),
                    r.get("customer_phone",""), self._bottle_name(r.get("bottle_code","")),
                    s.get("name",""), round(lab[0],2), round(lab[1],2), round(lab[2],2),
                    s.get("total_grams",0), r.get("updated_at",""),
                ])

        widths = [13, 26, 16, 14, 16, 14, 9, 9, 9, 11, 17]
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wdt

        try:
            wb.save(path)
        except Exception as e:
            QMessageBox.warning(self, "导出失败", f"保存文件时出错：{e}")
            return
        QMessageBox.information(self, "导出成功", f"已导出 {len(records)} 个配色模板到：\n{path}")
