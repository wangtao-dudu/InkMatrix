# -*- coding: utf-8 -*-
"""
finance_page.py  (UI层 - 财务管理)
================================================
回答一个问题：**这个客户到底该付我多少钱、付了多少、还欠多少、欠了多久。**

── 数据从哪来 ────────────────────────────────────────────────
不重新录一遍。财务【挂靠在进销存的进货记录上】：

    进销存记事实：客户A 7/14 送来 1000 个 100ml 白瓶要印
    财务算钱    ：1000 × 2.5 = 2500，优惠 200 → 实际应收 2300
                  客户先付了 1000 → 还欠 1300，已欠 6 天

数量/产品/客户/日期这些全部实时从进销存读，这边一份都不存——存两份
迟早对不上。财务只管进销存管不了的三件事：**单价、优惠、收款流水**。

── 为什么收款要按"笔"记，而不是只存一个"已付金额" ──────────
客户很少一次付清。今天付 1000，下月付 800，再下月付清。只存总数，就
永远说不清这 1800 是什么时候、怎么付的——对账时没有任何凭据可查。
所以按笔记流水，"已付总额"是算出来的。

── 为什么默认只看"进货" ──────────────────────────────────────
对丝印厂来说，客户把瓶子送进来（进货）就是活儿的开始，报价按这批的数量
算。出货是把印完的成品送回去，不是第二次收钱。所以应收账款挂在进货上。
（如果你的业务是按出货结算，右上角可以切换。）
"""
import datetime
from typing import List, Dict, Optional

from PyQt6.QtCore import Qt, pyqtSignal
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QComboBox,
    QTableWidget, QTableWidgetItem, QHeaderView, QMessageBox, QDialog,
    QDoubleSpinBox, QFormLayout, QDateEdit, QDialogButtonBox, QFileDialog,
    QTextEdit, QAbstractItemView,
)
from PyQt6.QtCore import QDate

from dal_layer.data_manager import FinanceRepository
from ui_layer.widgets import Card, PrimaryButton, GhostButton, DangerButton, StatCard, ResponsiveRow

# 账龄分级：欠得越久越该催，颜色越重
_AGING_BANDS = [
    (30,  "#f0fdf4", "#166534"),   # 30天内：绿
    (60,  "#fefce8", "#854d0e"),   # 30-60：黄
    (90,  "#fff7ed", "#9a3412"),   # 60-90：橙
    (999, "#fef2f2", "#991b1b"),   # 90天以上：红
]


def _aging_color(days: Optional[int]):
    if days is None:
        return None, None
    for limit, bg, fg in _AGING_BANDS:
        if days <= limit:
            return bg, fg
    return _AGING_BANDS[-1][1], _AGING_BANDS[-1][2]


class PaymentDialog(QDialog):
    """记一笔收款。"""

    def __init__(self, parent, customer: str, product: str, balance: float):
        super().__init__(parent)
        self.setWindowTitle("记录收款")
        self.setMinimumWidth(420)
        root = QVBoxLayout(self)

        info = QLabel(
            f"客户：<b>{customer}</b>　产品：<b>{product}</b><br>"
            f"当前尚欠：<b style='color:#dc2626;'>¥{balance:.2f}</b>"
        )
        info.setStyleSheet("background:#f8fafc; border:1px solid #e2e8f0; "
                           "border-radius:8px; padding:10px 12px;")
        root.addWidget(info)

        form = QFormLayout()
        self.amount_spin = QDoubleSpinBox()
        self.amount_spin.setRange(0.01, 100000000)
        self.amount_spin.setDecimals(2)
        self.amount_spin.setValue(max(0.01, balance))   # 默认填满欠款，一次结清最常见
        self.amount_spin.setPrefix("¥ ")

        self.date_edit = QDateEdit(QDate.currentDate())
        self.date_edit.setCalendarPopup(True)
        self.date_edit.setDisplayFormat("yyyy-MM-dd")

        self.method_combo = QComboBox()
        self.method_combo.addItems(["银行转账", "微信", "支付宝", "现金", "承兑汇票", "其他"])

        self.note_edit = QLineEdit()
        self.note_edit.setPlaceholderText("选填，例如：对方财务备注的流水号")

        form.addRow("收款金额：", self.amount_spin)
        form.addRow("收款日期：", self.date_edit)
        form.addRow("收款方式：", self.method_combo)
        form.addRow("备注：", self.note_edit)
        root.addLayout(form)

        hint = QLabel("💡 支持分期收款——每付一笔就记一笔，系统自动累计已付总额和尚欠金额。")
        hint.setWordWrap(True)
        hint.setStyleSheet("color:#64748b; font-size:11px;")
        root.addWidget(hint)

        btns = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        root.addWidget(btns)

    def get_payment(self) -> Dict:
        return {
            "amount": self.amount_spin.value(),
            "date": self.date_edit.date().toString("yyyy-MM-dd"),
            "method": self.method_combo.currentText(),
            "note": self.note_edit.text().strip(),
        }


class PriceDialog(QDialog):
    """改单价 / 改优惠。"""

    def __init__(self, parent, inv_record: Dict, calc: Dict):
        super().__init__(parent)
        self.setWindowTitle("调整单价与优惠")
        self.setMinimumWidth(440)
        root = QVBoxLayout(self)

        qty = calc["quantity"]
        info = QLabel(
            f"客户：<b>{inv_record.get('customer','')}</b>　"
            f"产品：<b>{inv_record.get('product_name','')}</b><br>"
            f"数量：<b>{qty:g}</b>（来自进销存，此处不可改；数量错了请去【进销存管理】改那条记录）"
        )
        info.setWordWrap(True)
        info.setStyleSheet("background:#f8fafc; border:1px solid #e2e8f0; "
                           "border-radius:8px; padding:10px 12px;")
        root.addWidget(info)

        form = QFormLayout()
        self.price_spin = QDoubleSpinBox()
        self.price_spin.setRange(0, 1000000)
        self.price_spin.setDecimals(2)
        self.price_spin.setPrefix("¥ ")
        self.price_spin.setValue(calc["unit_price"])
        self.price_spin.valueChanged.connect(self._recalc)

        self.discount_spin = QDoubleSpinBox()
        self.discount_spin.setRange(0, 100000000)
        self.discount_spin.setDecimals(2)
        self.discount_spin.setPrefix("¥ ")
        self.discount_spin.setValue(calc["discount"])
        self.discount_spin.valueChanged.connect(self._recalc)

        form.addRow("成交单价：", self.price_spin)
        form.addRow("优惠金额：", self.discount_spin)
        root.addLayout(form)

        self._qty = qty
        self.result_label = QLabel()
        self.result_label.setStyleSheet(
            "background:#ecfdf5; border:1px solid #a7f3d0; border-radius:8px; "
            "padding:10px 12px; font-size:13px;"
        )
        root.addWidget(self.result_label)
        self._recalc()

        btns = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        root.addWidget(btns)

    def _recalc(self):
        gross = self._qty * self.price_spin.value()
        recv = gross - self.discount_spin.value()
        self.result_label.setText(
            f"合计 = {self._qty:g} × ¥{self.price_spin.value():.2f} = <b>¥{gross:.2f}</b><br>"
            f"实际应收 = ¥{gross:.2f} − ¥{self.discount_spin.value():.2f} = "
            f"<b style='color:#059669; font-size:15px;'>¥{recv:.2f}</b>"
        )

    def get_values(self):
        return self.price_spin.value(), self.discount_spin.value()


class FinancePage(QWidget):

    status_message = pyqtSignal(str)

    def __init__(self, inventory_repo, finance_repo, company_repo=None, parent=None):
        super().__init__(parent)
        self.inventory_repo = inventory_repo
        self.finance_repo = finance_repo
        self.company_repo = company_repo
        self._rows: List[Dict] = []     # 当前表格每一行对应的完整数据
        self._build_ui()
        self.refresh_all()

    # ------------------------------------------------------------------

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 20, 24, 20)
        root.setSpacing(14)

        title = QLabel("财务管理 · 应收账款")
        title.setObjectName("PageTitle")
        subtitle = QLabel(
            "数据来自【进销存管理】的进货记录——数量和产品在那边录，这边只管单价、优惠和收款。"
            "两边永远是同一份数据，不会对不上"
        )
        subtitle.setObjectName("PageSubtitle")
        root.addWidget(title)
        root.addWidget(subtitle)

        stats = QHBoxLayout()
        self.stat_receivable = StatCard("¥0", "实际应收合计", accent_hex="#0d9488")
        self.stat_paid = StatCard("¥0", "已收合计", accent_hex="#16a34a")
        self.stat_balance = StatCard("¥0", "尚欠合计", accent_hex="#dc2626")
        self.stat_overdue = StatCard("0", "超60天未结清", accent_hex="#b45309")
        for s in (self.stat_receivable, self.stat_paid, self.stat_balance, self.stat_overdue):
            stats.addWidget(s)
        root.addLayout(stats)

        # 孤儿财务记录警告（平时隐藏）
        self.orphan_banner = QWidget()
        ob = QHBoxLayout(self.orphan_banner)
        ob.setContentsMargins(12, 10, 12, 10)
        self.orphan_label = QLabel("")
        self.orphan_label.setWordWrap(True)
        self.orphan_label.setStyleSheet("color:#7f1d1d; font-size:12px; background:transparent;")
        ob.addWidget(self.orphan_label, 1)
        self.orphan_btn = DangerButton("🧹 处理孤儿记录")
        self.orphan_btn.setMinimumHeight(32)
        self.orphan_btn.clicked.connect(self._on_purge_orphans)
        ob.addWidget(self.orphan_btn)
        self.orphan_banner.setStyleSheet(
            "background:#fef2f2; border:1px solid #fecaca; border-radius:8px;"
        )
        self.orphan_banner.setVisible(False)
        root.addWidget(self.orphan_banner)

        card = Card("应收账款明细")
        lay = card.body_layout

        bar = QHBoxLayout()
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("搜索客户 / 产品 / 单号…")
        self.search_edit.textChanged.connect(self.refresh_all)
        bar.addWidget(self.search_edit, 2)

        self.status_filter = QComboBox()
        self.status_filter.addItem("全部状态", None)
        for s in ("未付款", "部分已付", "已结清"):
            self.status_filter.addItem(s, s)
        self.status_filter.currentIndexChanged.connect(self.refresh_all)
        bar.addWidget(self.status_filter)

        self.type_filter = QComboBox()
        self.type_filter.addItem("按进货结算（默认）", "进货")
        self.type_filter.addItem("按出货结算", "出货")
        self.type_filter.setToolTip(
            "丝印厂通常在客户把瓶子送进来（进货）时就按数量报价，所以应收挂在进货上。\n"
            "如果你的业务是按出货结算，切到这一项。"
        )
        self.type_filter.currentIndexChanged.connect(self.refresh_all)
        bar.addWidget(self.type_filter)

        bar.addStretch(1)
        price_btn = PrimaryButton("💰 改单价/优惠")
        price_btn.clicked.connect(self._on_edit_price)
        pay_btn = PrimaryButton("✅ 记录收款")
        pay_btn.clicked.connect(self._on_add_payment)
        hist_btn = GhostButton("📜 收款流水")
        hist_btn.clicked.connect(self._on_view_payments)
        export_btn = GhostButton("📊 导出Excel")
        export_btn.clicked.connect(self._on_export)
        for b in (price_btn, pay_btn, hist_btn, export_btn):
            bar.addWidget(b)
        lay.addLayout(bar)

        self.table = QTableWidget(0, 12)
        self.table.setHorizontalHeaderLabels([
            "日期", "客户", "产品名称", "单号", "数量",
            "单价(元)", "合计(元)", "优惠(元)", "实际应收(元)",
            "已付(元)", "尚欠(元)", "状态 / 账龄",
        ])
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setDefaultSectionSize(34)
        self.table.doubleClicked.connect(self._on_edit_price)
        lay.addWidget(self.table, 1)

        hint = QLabel(
            "💡 双击任意一行可以直接改单价/优惠。数量或产品名录错了，请去【进销存管理】改那条记录——"
            "财务这边会自动跟着变，不用改两遍。"
        )
        hint.setWordWrap(True)
        hint.setStyleSheet("color:#64748b; font-size:11px;")
        lay.addWidget(hint)

        root.addWidget(card, 1)

    # ------------------------------------------------------------------

    def _build_rows(self) -> List[Dict]:
        """把进销存记录 + 财务记录 合成一张应收账款表。"""
        want_type = self.type_filter.currentData()
        keyword = self.search_edit.text().strip().lower()
        want_status = self.status_filter.currentData()

        rows = []
        for inv in self.inventory_repo.load_all():
            if inv.get("movement_type") != want_type:
                continue
            if keyword:
                blob = " ".join([
                    inv.get("customer", "") or "", inv.get("product_name", "") or "",
                    inv.get("order_no", "") or "",
                ]).lower()
                if keyword not in blob:
                    continue

            fin = self.finance_repo.get_by_inventory_id(inv["id"])
            calc = FinanceRepository.compute(inv, fin)
            if want_status and calc["status"] != want_status:
                continue
            calc["aging"] = FinanceRepository.aging_days(inv.get("date", ""), calc["balance"])
            rows.append({"inv": inv, "fin": fin, "calc": calc})

        rows.sort(key=lambda r: (r["inv"].get("date", ""), r["inv"].get("created_at", "")), reverse=True)
        return rows

    def refresh_all(self):
        self._check_orphans()
        self._rows = self._build_rows()
        self.table.setRowCount(len(self._rows))

        tot_recv = tot_paid = tot_bal = 0.0
        overdue = 0
        overpaid_count = 0

        for i, row in enumerate(self._rows):
            inv, c = row["inv"], row["calc"]
            tot_recv += c["receivable"]
            tot_paid += c["paid"]
            tot_bal += c["balance"]
            if c["aging"] is not None and c["aging"] > 60:
                overdue += 1
            if c["balance"] < -0.001:
                overpaid_count += 1

            # 尚欠是负数 = 收多了。这不是"已结清"，是一笔要退给客户/要冲抵下一单的钱，
            # 必须单独标出来。不然它混在一堆"已结清"里，谁也不会发现多收了。
            overpaid = c["balance"] < -0.001
            if overpaid:
                aging_txt = f"⚠️ 多收 ¥{-c['balance']:.2f}"
            else:
                aging_txt = c["status"]
                if c["aging"] is not None:
                    aging_txt += f" · 欠{c['aging']}天"

            cells = [
                inv.get("date", ""), inv.get("customer", ""), inv.get("product_name", ""),
                inv.get("order_no", "") or "-", f"{c['quantity']:g}",
                f"{c['unit_price']:.2f}", f"{c['gross']:.2f}", f"{c['discount']:.2f}",
                f"{c['receivable']:.2f}", f"{c['paid']:.2f}", f"{c['balance']:.2f}",
                aging_txt,
            ]
            bg, fg = _aging_color(c["aging"])
            for col, text in enumerate(cells):
                item = QTableWidgetItem(text)
                if col == 0:
                    item.setData(Qt.ItemDataRole.UserRole, inv["id"])
                if col not in (1, 2):
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if col == 10 and c["balance"] > 0.001:
                    item.setForeground(QColor("#dc2626"))
                    f = item.font(); f.setBold(True); item.setFont(f)
                if col == 10 and overpaid:
                    item.setForeground(QColor("#7c3aed"))
                    f = item.font(); f.setBold(True); item.setFont(f)
                    item.setToolTip("收款金额超过了实际应收——多收的部分要退还客户或冲抵下一单")
                if col == 11:
                    if overpaid:
                        item.setBackground(QColor("#f5f3ff"))
                        item.setForeground(QColor("#6d28d9"))
                    elif c["status"] == "已结清":
                        item.setBackground(QColor("#f0fdf4"))
                        item.setForeground(QColor("#166534"))
                    elif bg:
                        item.setBackground(QColor(bg))
                        item.setForeground(QColor(fg))
                self.table.setItem(i, col, item)

        self.stat_receivable.set_value(f"¥{tot_recv:,.2f}")
        self.stat_paid.set_value(f"¥{tot_paid:,.2f}")
        self.stat_balance.set_value(f"¥{tot_bal:,.2f}")
        # 有多收款的时候，把"超60天"这张卡临时借来报警——多收款比逾期更需要马上处理
        if overpaid_count:
            self.stat_overdue.set_value(f"{overdue} / ⚠️{overpaid_count}笔多收")
        else:
            self.stat_overdue.set_value(str(overdue))

    # ------------------------------------------------------------------

    def _check_orphans(self):
        """
        检查有没有"孤儿"财务记录——它挂的那条进销存记录已经不在了。

        这种记录在表格里是【看不见】的（表格是从进销存反查出来的），所以如果
        不主动提醒，里面挂着的收款金额就永远查不到、也删不掉。必须显式报出来。
        """
        orphans = self.finance_repo.find_orphans(self.inventory_repo)
        if not orphans:
            self.orphan_banner.setVisible(False)
            return

        total_paid = sum(FinanceRepository.paid_total(o) for o in orphans)
        money_part = ""
        if total_paid > 0.001:
            money_part = (
                f"，其中<b>挂着已收款 ¥{total_paid:.2f}</b>——"
                f"这些钱是真收过的，但对应的进货记录已经不在了"
            )
        self.orphan_label.setText(
            f"⚠️ <b>发现 {len(orphans)} 条孤儿财务记录</b>{money_part}。<br>"
            f"它们指向的进销存记录已被删除，所以在下面的表格里看不到，但一直占着账。"
            f"点右边处理。"
        )
        self.orphan_banner.setVisible(True)

    def _on_purge_orphans(self):
        orphans = self.finance_repo.find_orphans(self.inventory_repo)
        if not orphans:
            QMessageBox.information(self, "提示", "现在没有孤儿记录。")
            return

        total_paid = sum(FinanceRepository.paid_total(o) for o in orphans)

        lines = [f"共 {len(orphans)} 条孤儿财务记录：", ""]
        for o in orphans[:10]:
            paid = FinanceRepository.paid_total(o)
            lines.append(
                f"  · 关联进货单 {o.get('inventory_id','?')}　"
                f"单价 ¥{o.get('unit_price',0):.2f}　优惠 ¥{o.get('discount',0):.2f}　"
                f"已收 ¥{paid:.2f}"
            )
        if len(orphans) > 10:
            lines.append(f"  …还有 {len(orphans)-10} 条")

        warn = ""
        if total_paid > 0.001:
            warn = (
                f"\n\n🚨 这些记录里挂着【已收款合计 ¥{total_paid:.2f}】。\n"
                f"这笔钱是真的收过的——很可能是有人误删了对应的进货记录。\n\n"
                f"删掉这些孤儿记录，等于承认这笔收款在账上永久消失。\n"
                f"建议先把这些金额抄下来，去【进销存管理】把对应的进货记录补录回来，"
                f"再重新记一次收款。"
            )
        else:
            warn = "\n\n这些记录里没有任何收款，只是设过单价/优惠，删掉是安全的。"

        box = QMessageBox(self)
        box.setWindowTitle("处理孤儿财务记录")
        box.setIcon(QMessageBox.Icon.Warning if total_paid > 0.001 else QMessageBox.Icon.Question)
        box.setText("\n".join(lines) + warn + "\n\n确定要删除这些孤儿记录吗？")
        box.setStandardButtons(QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        box.setDefaultButton(QMessageBox.StandardButton.No)
        if box.exec() != QMessageBox.StandardButton.Yes:
            return

        n = self.finance_repo.purge_orphans(self.inventory_repo)
        self.refresh_all()
        self.status_message.emit(f"已清理 {n} 条孤儿财务记录。")

    def _selected(self) -> Optional[Dict]:
        row = self.table.currentRow()
        if row < 0 or row >= len(self._rows):
            QMessageBox.information(self, "提示", "请先在表格中点击选中一行。")
            return None
        return self._rows[row]

    def _on_edit_price(self):
        row = self._selected()
        if not row:
            return
        inv, fin, calc = row["inv"], row["fin"], row["calc"]
        dlg = PriceDialog(self, inv, calc)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        price, discount = dlg.get_values()
        rec = fin or {"inventory_id": inv["id"], "payments": []}
        rec["unit_price"] = price
        rec["discount"] = discount
        self.finance_repo.upsert(rec)
        self.refresh_all()
        self.status_message.emit(
            f"已更新 [{inv.get('customer','')} · {inv.get('product_name','')}] "
            f"的单价 ¥{price:.2f} / 优惠 ¥{discount:.2f}"
        )

    def _on_add_payment(self):
        row = self._selected()
        if not row:
            return
        inv, calc = row["inv"], row["calc"]
        if calc["receivable"] <= 0.001:
            QMessageBox.information(
                self, "提示",
                "这条记录的应收金额是 0，请先用【💰 改单价/优惠】把成交单价填上。"
            )
            return
        if calc["balance"] <= 0.001:
            confirm = QMessageBox.question(
                self, "已结清",
                "这笔账已经结清了，还要再记一笔收款吗？（多收的部分会显示为负数欠款）"
            )
            if confirm != QMessageBox.StandardButton.Yes:
                return

        dlg = PaymentDialog(self, inv.get("customer", ""), inv.get("product_name", ""), calc["balance"])
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        p = dlg.get_payment()
        self.finance_repo.add_payment(inv["id"], p["amount"], p["date"], p["method"], p["note"])

        fin = self.finance_repo.get_by_inventory_id(inv["id"])
        new_calc = FinanceRepository.compute(inv, fin)
        self.refresh_all()

        tail = "✅ 该笔账款已全部结清。" if new_calc["balance"] <= 0.001 \
            else f"仍欠 ¥{new_calc['balance']:.2f}。"
        self.status_message.emit(f"已收款 ¥{p['amount']:.2f}（{p['method']}）。{tail}")

    def _on_view_payments(self):
        row = self._selected()
        if not row:
            return
        inv, fin, calc = row["inv"], row["fin"], row["calc"]
        payments = (fin or {}).get("payments", [])
        if not payments:
            QMessageBox.information(self, "收款流水", "这笔账还没有任何收款记录。")
            return

        lines = [f"客户：{inv.get('customer','')}　产品：{inv.get('product_name','')}",
                 f"实际应收：¥{calc['receivable']:.2f}", ""]
        for i, p in enumerate(payments, 1):
            note = f"　{p.get('note','')}" if p.get("note") else ""
            lines.append(f"{i}. {p.get('date','')}　¥{p.get('amount',0):.2f}　"
                         f"{p.get('method','')}{note}")
        lines += ["", f"已付合计：¥{calc['paid']:.2f}",
                  f"尚欠：¥{calc['balance']:.2f}"]

        dlg = QDialog(self)
        dlg.setWindowTitle("收款流水")
        dlg.setMinimumSize(460, 340)
        v = QVBoxLayout(dlg)
        te = QTextEdit()
        te.setReadOnly(True)
        te.setPlainText("\n".join(lines))
        v.addWidget(te)
        b = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        b.rejected.connect(dlg.reject)
        b.accepted.connect(dlg.accept)
        v.addWidget(b)
        dlg.exec()

    def _on_export(self):
        if not self._rows:
            QMessageBox.information(self, "提示", "当前没有可导出的数据。")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "导出应收账款",
            f"应收账款_{datetime.date.today().isoformat()}.xlsx",
            "Excel 文件 (*.xlsx)"
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            QMessageBox.warning(self, "缺少依赖", "导出Excel需要 openpyxl，请先安装：pip install openpyxl")
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "应收账款"
        headers = ["日期", "客户", "产品名称", "单号", "数量", "单价(元)", "合计(元)",
                   "优惠(元)", "实际应收(元)", "已付(元)", "尚欠(元)", "状态", "账龄(天)"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0D9488")
            c.alignment = Alignment(horizontal="center")

        for row in self._rows:
            inv, c = row["inv"], row["calc"]
            ws.append([
                inv.get("date", ""), inv.get("customer", ""), inv.get("product_name", ""),
                inv.get("order_no", ""), c["quantity"], c["unit_price"], c["gross"],
                c["discount"], c["receivable"], c["paid"], c["balance"],
                c["status"], c["aging"] if c["aging"] is not None else "",
            ])

        widths = [12, 18, 26, 14, 10, 11, 12, 11, 14, 12, 12, 11, 10]
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wdt

        try:
            wb.save(path)
        except Exception as e:
            QMessageBox.warning(self, "导出失败", f"保存文件时出错：{e}")
            return
        QMessageBox.information(self, "导出成功", f"已导出 {len(self._rows)} 条应收账款到：\n{path}")
