# -*- coding: utf-8 -*-
"""
ink_stock_page.py  (UI层 - 油墨库存管理页面，独立页面)
============================================================
跟【生产追踪】【进销存管理】完全独立——那两个是"工厂做的成品/送给
客户的货"，这个页面管的是【油墨这个原料/商品本身】的库存、采购、
销售，服务的是"我自己就是油墨供应商，要盘点自己的库存和销售情况"
这个视角。

功能：
    - 进出库记录（采购入库 / 销售出库），带价格、供应商/客户、备注
    - 库存盘点：录入实际清点数量，系统自动算出差异并调整库存
    - 销售统计总览：按客户汇总卖了多少、多少钱
    - 搜索、导出Excel
"""

import datetime
from typing import List, Dict

from PyQt6.QtCore import Qt, QDate, pyqtSignal
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QComboBox,
    QDoubleSpinBox, QDateEdit, QTextEdit, QMessageBox, QFileDialog,
    QTableWidget, QTableWidgetItem, QHeaderView, QScrollArea, QTabWidget
)

from ui_layer.widgets import Card, PrimaryButton, GhostButton, DangerButton, StatCard, ResponsiveRow


class InkStockPage(QWidget):

    ink_db_changed = pyqtSignal()

    def __init__(self, ink_repo, ink_stock_repo, parent=None):
        super().__init__(parent)
        self.ink_repo = ink_repo
        self.ink_stock_repo = ink_stock_repo
        self._build_ui()
        self.refresh_all()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 20, 24, 20)
        root.setSpacing(14)

        title = QLabel("油墨库存管理")
        title.setObjectName("PageTitle")
        subtitle = QLabel("油墨原料本身的采购入库、销售出库、库存盘点——独立于成品进销存，专供油墨供应商视角")
        subtitle.setObjectName("PageSubtitle")
        root.addWidget(title)
        root.addWidget(subtitle)

        stat_row = QHBoxLayout()
        stat_row.setSpacing(12)
        self.stat_purchase = StatCard("¥0", "累计采购金额", accent_hex="#2563eb")
        self.stat_sales = StatCard("¥0", "累计销售金额", accent_hex="#0d9488")
        self.stat_stock_value = StatCard("¥0", "当前库存估值", accent_hex="#d97706")
        self.stat_low_stock = StatCard("0", "低库存/缺货油墨数", accent_hex="#dc2626")
        for c in (self.stat_purchase, self.stat_sales, self.stat_stock_value, self.stat_low_stock):
            stat_row.addWidget(c)
        stat_row.addStretch(1)
        root.addLayout(stat_row)

        tabs = QTabWidget()
        # "当前库存"放在第一个：这是【问题3】要的"账实统一"最直观的落地——
        # 这张表和【油墨库管理】的库存列读的是同一份 inks_db.json 里的
        # stock_ml，永远是同一个数字，不可能对不上。
        tabs.addTab(self._build_stock_overview_tab(), "🛢️ 当前库存")
        tabs.addTab(self._wrap_scrollable(self._build_ledger_tab()), "📋 进出库记录")
        tabs.addTab(self._wrap_scrollable(self._build_stocktake_tab()), "🔍 库存盘点")
        tabs.addTab(self._build_sales_summary_tab(), "📊 客户销售统计")
        root.addWidget(tabs, 1)

    # ================================================================
    # ⓪ 当前库存总览（唯一权威库存视图）
    # ================================================================

    def _build_stock_overview_tab(self) -> QWidget:
        container = QWidget()
        lay = QVBoxLayout(container)

        banner = QLabel(
            "✅ 这里显示的库存数量，和【油墨库管理】页看到的是同一个数字——"
            "两处读的都是同一份油墨档案里的库存字段，不存在两套账。\n"
            "🔒 库存只能通过本页的【采购入库】/【销售出库】/【库存盘点】改动；"
            "在【油墨库管理】那边编辑油墨档案（改名字、改光谱）不会动到库存分毫。"
        )
        banner.setWordWrap(True)
        banner.setStyleSheet(
            "color:#065f46; background:#ecfdf5; border:1px solid #a7f3d0; "
            "border-radius:8px; padding:10px 12px; font-size:12px;"
        )
        lay.addWidget(banner)

        card = Card("全部油墨当前库存")
        clay = card.body_layout

        # 库存按 ml 记（采购是按桶/升买的），配墨按 g 称（天平称的是重量），
        # 两个单位都列出来，并且把密度也摆在旁边，换算关系一目了然（问题2）。
        self.overview_table = QTableWidget(0, 7)
        self.overview_table.setHorizontalHeaderLabels([
            "编号", "名称", "库存(ml)", "折合重量(g)", "密度(g/ml)", "预警线(ml)", "状态"
        ])
        hdr = self.overview_table.horizontalHeader()
        hdr.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.overview_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.overview_table.setAlternatingRowColors(True)
        clay.addWidget(self.overview_table, 1)

        # 低库存预警线的设置入口挪到这里——因为它属于"库存管理"的范畴，
        # 跟油墨的光谱/密度那些档案属性不是一回事，放在这一页才合逻辑。
        thr_row = QHBoxLayout()
        thr_row.addWidget(QLabel("选中油墨的低库存预警线(ml)："))
        self.threshold_spin = QDoubleSpinBox()
        self.threshold_spin.setRange(0, 1000000)
        self.threshold_spin.setDecimals(0)
        self.threshold_spin.setValue(100)
        thr_row.addWidget(self.threshold_spin)
        set_thr_btn = PrimaryButton("💾 设置预警线")
        set_thr_btn.clicked.connect(self._on_set_threshold)
        thr_row.addWidget(set_thr_btn)
        thr_row.addStretch(1)
        clay.addLayout(thr_row)

        lay.addWidget(card, 1)
        return container

    def _refresh_stock_overview(self):
        records = self.ink_repo.load_all()
        self.overview_table.setRowCount(len(records))
        for row, r in enumerate(records):
            stock_ml = r.get("stock_ml", 0.0) or 0.0
            density = r.get("density", 1.0) or 1.0
            grams = self.ink_repo.ml_to_grams(r, stock_ml)
            threshold = r.get("low_stock_threshold_ml", 100.0) or 0.0
            status = self.ink_repo.stock_status(r)

            code_item = QTableWidgetItem(r["code"])
            code_item.setData(Qt.ItemDataRole.UserRole, r["code"])
            items = [
                code_item,
                QTableWidgetItem(r["name"]),
                QTableWidgetItem(f"{stock_ml:g}"),
                QTableWidgetItem(f"{grams:.0f}"),
                QTableWidgetItem(f"{density:.2f}"),
                QTableWidgetItem(f"{threshold:g}"),
                QTableWidgetItem(status),
            ]
            for col, item in enumerate(items):
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if status == "缺货":
                    item.setBackground(QColor("#fee2e2"))
                elif status == "库存不足":
                    item.setBackground(QColor("#fef9c3"))
                self.overview_table.setItem(row, col, item)

    def _on_set_threshold(self):
        row = self.overview_table.currentRow()
        if row < 0:
            QMessageBox.information(self, "提示", "请先在表格中点击选中一支油墨。")
            return
        code = self.overview_table.item(row, 0).data(Qt.ItemDataRole.UserRole)
        self.ink_repo.set_low_stock_threshold(code, self.threshold_spin.value())
        self.refresh_all()
        self.ink_db_changed.emit()
        QMessageBox.information(
            self, "已设置",
            f"[{code}] 的低库存预警线已设为 {self.threshold_spin.value():g} ml，"
            f"库存降到这个数以下时会自动标黄提醒。"
        )

    @staticmethod
    def _wrap_scrollable(inner: QWidget) -> QWidget:
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        scroll.setWidget(inner)
        return scroll

    # ================================================================
    # ① 进出库记录
    # ================================================================

    def _build_ledger_tab(self) -> QWidget:
        body = ResponsiveRow(breakpoint=900)

        new_card = Card("新增进出库记录")
        nlay = new_card.body_layout
        nlay.addWidget(QLabel("类型："))
        self.stock_type_combo = QComboBox()
        self.stock_type_combo.addItems(["采购入库", "销售出库"])
        self.stock_type_combo.currentIndexChanged.connect(self._on_stock_type_changed)
        nlay.addWidget(self.stock_type_combo)

        nlay.addWidget(QLabel("日期："))
        self.stock_date_edit = QDateEdit()
        self.stock_date_edit.setCalendarPopup(True)
        self.stock_date_edit.setDisplayFormat("yyyy-MM-dd")
        self.stock_date_edit.setDate(QDate.currentDate())
        nlay.addWidget(self.stock_date_edit)

        nlay.addWidget(QLabel("油墨（选已有的，或直接输入新油墨名称自动建档）："))
        self.stock_ink_combo = QComboBox()
        self.stock_ink_combo.setEditable(True)
        nlay.addWidget(self.stock_ink_combo)

        row1 = QHBoxLayout()
        qty_col = QVBoxLayout()
        qty_col.addWidget(QLabel("数量(ml)："))
        self.stock_qty_spin = QDoubleSpinBox()
        self.stock_qty_spin.setRange(0, 1000000)
        self.stock_qty_spin.setDecimals(0)
        self.stock_qty_spin.valueChanged.connect(self._update_stock_amount_preview)
        qty_col.addWidget(self.stock_qty_spin)
        price_col = QVBoxLayout()
        price_col.addWidget(QLabel("单价(元/ml)："))
        self.stock_price_spin = QDoubleSpinBox()
        self.stock_price_spin.setRange(0, 100000)
        self.stock_price_spin.setDecimals(2)
        self.stock_price_spin.valueChanged.connect(self._update_stock_amount_preview)
        price_col.addWidget(self.stock_price_spin)
        row1.addLayout(qty_col)
        row1.addLayout(price_col)
        nlay.addLayout(row1)

        self.stock_amount_label = QLabel("金额：¥0.00")
        self.stock_amount_label.setStyleSheet("color:#0f172a; font-weight:600;")
        nlay.addWidget(self.stock_amount_label)

        self.stock_party_label = QLabel("供应商：")
        nlay.addWidget(self.stock_party_label)
        self.stock_party_edit = QLineEdit()
        nlay.addWidget(self.stock_party_edit)

        nlay.addWidget(QLabel("备注："))
        self.stock_notes_edit = QTextEdit()
        self.stock_notes_edit.setMaximumHeight(60)
        nlay.addWidget(self.stock_notes_edit)

        save_btn = PrimaryButton("💾 保存记录")
        save_btn.clicked.connect(self._on_save_stock_record)
        nlay.addWidget(save_btn)
        body.add_panel(new_card, 1)

        history_card = Card("进出库历史")
        hlay = history_card.body_layout
        toolbar = QHBoxLayout()
        self.stock_search_edit = QLineEdit()
        self.stock_search_edit.setPlaceholderText("🔍 按油墨名称/供应商/客户搜索…")
        self.stock_search_edit.textChanged.connect(self._refresh_stock_table)
        toolbar.addWidget(self.stock_search_edit)
        toolbar.addStretch(1)
        export_btn = GhostButton("📊 导出Excel")
        export_btn.clicked.connect(self._on_export_stock_excel)
        del_btn = DangerButton("－ 删除选中")
        del_btn.setMinimumHeight(34)
        del_btn.clicked.connect(self._on_delete_stock_record)
        toolbar.addWidget(export_btn)
        toolbar.addWidget(del_btn)
        hlay.addLayout(toolbar)

        self.stock_table = QTableWidget(0, 8)
        self.stock_table.setHorizontalHeaderLabels(
            ["日期", "油墨", "类型", "数量(ml)", "单价", "金额", "供应商/客户", "备注"]
        )
        self.stock_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.stock_table.verticalHeader().setDefaultSectionSize(32)
        hlay.addWidget(self.stock_table, 1)
        body.add_panel(history_card, 1)

        return body

    def _on_stock_type_changed(self, _index: int):
        is_in = self.stock_type_combo.currentText() == "采购入库"
        self.stock_party_label.setText("供应商：" if is_in else "客户：")
        self.stock_party_edit.setPlaceholderText("例如：某某原料厂" if is_in else "例如：客户A")

    def _update_stock_amount_preview(self):
        amount = self.stock_qty_spin.value() * self.stock_price_spin.value()
        self.stock_amount_label.setText(f"金额：¥{amount:,.2f}")

    def _refresh_stock_ink_combo(self):
        current = self.stock_ink_combo.currentData()
        self.stock_ink_combo.blockSignals(True)
        self.stock_ink_combo.clear()
        for r in self.ink_repo.load_all():
            self.stock_ink_combo.addItem(f"[{r['code']}] {r['name']}", r["code"])
        if current:
            idx = self.stock_ink_combo.findData(current)
            if idx >= 0:
                self.stock_ink_combo.setCurrentIndex(idx)
        self.stock_ink_combo.blockSignals(False)

    def _on_save_stock_record(self):
        ink_code = self.stock_ink_combo.currentData()
        typed_text = self.stock_ink_combo.currentText().strip()

        if not ink_code:
            # 下拉框里没有对应的已有油墨，说明用户是直接输入了一个新名字，
            # 自动在【油墨库管理】里建一条基础档案（灰阶占位光谱，之后
            # 建议去用色相环补上真实颜色），这样"入库"和"油墨档案"两件事
            # 就不用分两个页面各跑一次了。
            if not typed_text:
                QMessageBox.warning(self, "提示", "请选择已有油墨，或输入一个新油墨名称。")
                return
            new_code = f"NEW-{len(self.ink_repo.load_all()) + 1:03d}"
            self.ink_repo.upsert({
                "code": new_code, "name": typed_text, "category": "color",
                "K": [20.0] * 31, "S": [10.0] * 31,
                "density": 1.05, "dry_back_factor": 0.25,
                "stock_ml": 0.0, "low_stock_threshold_ml": 100.0,
            })
            ink_code = new_code
            self.ink_db_changed.emit()
            QMessageBox.information(
                self, "已自动建档",
                f"检测到【{typed_text}】还没有油墨档案，已自动在【油墨库管理】里创建一条基础记录"
                f"（先用灰阶占位光谱），建议之后去那边用色相环补上真实颜色。"
            )
            self._refresh_stock_ink_combo()
            idx = self.stock_ink_combo.findData(new_code)
            if idx >= 0:
                self.stock_ink_combo.setCurrentIndex(idx)

        if self.stock_qty_spin.value() <= 0:
            QMessageBox.warning(self, "提示", "数量需要大于0。")
            return

        ink_record = self.ink_repo.get_by_code(ink_code)
        ui_type = self.stock_type_combo.currentText()
        movement_type = "入库" if ui_type == "采购入库" else "出库"
        now = datetime.datetime.now()
        record = {
            "id": f"IS-{now.strftime('%Y%m%d%H%M%S%f')}",
            "date": self.stock_date_edit.date().toString("yyyy-MM-dd"),
            "movement_type": movement_type,
            "ink_code": ink_code,
            "ink_name": ink_record["name"] if ink_record else ink_code,
            "quantity_ml": self.stock_qty_spin.value(),
            "price": self.stock_price_spin.value(),
            "amount": round(self.stock_qty_spin.value() * self.stock_price_spin.value(), 2),
            "supplier": self.stock_party_edit.text().strip() if movement_type == "入库" else "",
            "customer": self.stock_party_edit.text().strip() if movement_type == "出库" else "",
            "notes": self.stock_notes_edit.toPlainText().strip(),
            "created_at": now.strftime("%Y-%m-%d %H:%M"),
        }
        self.ink_stock_repo.upsert(record)

        # 记完流水，立刻同步更新油墨档案里的库存字段（唯一真值源）
        delta = self.stock_qty_spin.value() if movement_type == "入库" else -self.stock_qty_spin.value()
        updated = self.ink_repo.adjust_stock(ink_code, delta)

        # ★【问题3修复】发出信号，让【油墨库管理】页立刻重新读盘刷新。
        # 以前漏了这一句，结果是：这边入库完库存已经变了，但用户切到
        # 【油墨库管理】看到的还是打开页面时缓存的旧数字，看上去就像
        # "两个页面的库存对不上"。现在改完立刻通知，两边永远同步。
        self.ink_db_changed.emit()

        # 顺手把换算后的克重也报出来（问题2）：入库按 ml 记，但师傅配墨
        # 时要按 g 称，直接告诉他这批墨能称出多少克，省得自己算。
        new_ml = (updated or {}).get("stock_ml", 0.0)
        new_g = self.ink_repo.ml_to_grams(updated or {}, new_ml)
        QMessageBox.information(
            self, "已保存",
            f"{ui_type}记录已保存。\n\n"
            f"【{record['ink_name']}】最新库存：{new_ml:g} ml  ≈  {new_g:.0f} g\n"
            f"（按密度 {(updated or {}).get('density', 1.0):.2f} g/ml 换算）\n\n"
            f"该数字已同步到【油墨库管理】页，两处显示完全一致。"
        )
        self.stock_qty_spin.setValue(0)
        self.stock_price_spin.setValue(0)
        self.stock_party_edit.clear()
        self.stock_notes_edit.clear()
        self.refresh_all()

    def _filtered_stock_records(self) -> List[Dict]:
        records = self.ink_stock_repo.load_all()
        keyword = self.stock_search_edit.text().strip().lower()
        if not keyword:
            return records
        return [
            r for r in records
            if keyword in (r.get("ink_name", "") or "").lower()
            or keyword in (r.get("supplier", "") or "").lower()
            or keyword in (r.get("customer", "") or "").lower()
        ]

    def _refresh_stock_table(self):
        records = self._filtered_stock_records()
        self.stock_table.setRowCount(len(records))
        for row, r in enumerate(records):
            type_item = QTableWidgetItem(r.get("movement_type", ""))
            type_item.setData(Qt.ItemDataRole.UserRole, r["id"])
            is_in = r.get("movement_type") == "入库"
            type_item.setForeground(QColor("#2563eb" if is_in else "#0d9488"))
            party = r.get("supplier") or r.get("customer") or "-"
            items = [
                QTableWidgetItem(r.get("date", "")),
                QTableWidgetItem(r.get("ink_name", "")),
                type_item,
                QTableWidgetItem(f"{r.get('quantity_ml', 0):g}"),
                QTableWidgetItem(f"{r.get('price', 0):.2f}"),
                QTableWidgetItem(f"{r.get('amount', 0):.2f}"),
                QTableWidgetItem(party),
                QTableWidgetItem(r.get("notes", "") or "-"),
            ]
            for col, item in enumerate(items):
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.stock_table.setItem(row, col, item)

    def _on_delete_stock_record(self):
        row = self.stock_table.currentRow()
        if row < 0:
            QMessageBox.information(self, "提示", "请先选中要删除的记录。")
            return
        record_id = self.stock_table.item(row, 2).data(Qt.ItemDataRole.UserRole)
        confirm = QMessageBox.question(self, "确认删除", "确定要删除这条进出库记录吗？（不会自动回滚已经扣减/增加的库存数量）")
        if confirm == QMessageBox.StandardButton.Yes:
            self.ink_stock_repo.delete(record_id)
            self._refresh_stock_table()

    def _on_export_stock_excel(self):
        records = self._filtered_stock_records()
        if not records:
            QMessageBox.information(self, "提示", "没有可导出的记录。")
            return
        default_name = f"油墨进出库记录_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "导出油墨进出库Excel", default_name, "Excel文件 (*.xlsx)")
        if not path:
            return
        try:
            self._write_ledger_excel(records, path)
        except Exception as exc:
            QMessageBox.warning(self, "导出失败", f"生成Excel文件时出错：{exc}")
            return
        QMessageBox.information(self, "导出成功", f"进出库记录已保存到：\n{path}")

    def _write_ledger_excel(self, records: List[Dict], path: str):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "油墨进出库记录"
        headers = ["日期", "油墨", "类型", "数量(ml)", "单价(元)", "金额(元)", "供应商/客户", "备注"]
        header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row_idx, r in enumerate(records, start=2):
            party = r.get("supplier") or r.get("customer") or ""
            values = [r.get("date", ""), r.get("ink_name", ""), r.get("movement_type", ""),
                      r.get("quantity_ml", 0), r.get("price", 0), r.get("amount", 0), party, r.get("notes", "")]
            for col, v in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col, value=v)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16
        wb.save(path)

    # ================================================================
    # ② 库存盘点
    # ================================================================

    def _build_stocktake_tab(self) -> QWidget:
        container = QWidget()
        lay = QVBoxLayout(container)

        hint = QLabel(
            "💡 库存盘点：实际去仓库数一下某支墨还剩多少毫升，填在下面\"实际数量\"，"
            "系统会自动跟账面库存对比算出差异，点\"确认调整\"后会把系统库存改成你数出来的实际数量，"
            "并生成一条\"盘点调整\"记录留痕（方便对账时知道哪次调整过、调了多少）。"
        )
        hint.setWordWrap(True)
        hint.setStyleSheet("color:#64748b; font-size:11px;")
        lay.addWidget(hint)

        form_card = Card("盘点单支油墨")
        flay = form_card.body_layout
        row = QHBoxLayout()
        row.addWidget(QLabel("油墨："))
        self.take_ink_combo = QComboBox()
        row.addWidget(self.take_ink_combo, 1)
        flay.addLayout(row)

        row2 = QHBoxLayout()
        row2.addWidget(QLabel("账面库存(ml)："))
        self.take_system_label = QLabel("--")
        self.take_system_label.setStyleSheet("font-weight:700;")
        row2.addWidget(self.take_system_label)
        row2.addWidget(QLabel("实际清点数量(ml)："))
        self.take_actual_spin = QDoubleSpinBox()
        self.take_actual_spin.setRange(0, 1000000)
        self.take_actual_spin.setDecimals(0)
        self.take_actual_spin.valueChanged.connect(self._update_stocktake_diff)
        row2.addWidget(self.take_actual_spin)
        flay.addLayout(row2)

        self.take_ink_combo.currentIndexChanged.connect(self._on_stocktake_ink_changed)

        self.take_diff_label = QLabel("差异：--")
        self.take_diff_label.setStyleSheet("font-weight:700; font-size:14px;")
        flay.addWidget(self.take_diff_label)

        confirm_btn = PrimaryButton("✅ 确认调整为实际数量")
        confirm_btn.clicked.connect(self._on_confirm_stocktake)
        flay.addWidget(confirm_btn)

        lay.addWidget(form_card)
        lay.addStretch(1)
        return container

    def _refresh_stocktake_ink_combo(self):
        current = self.take_ink_combo.currentData()
        self.take_ink_combo.blockSignals(True)
        self.take_ink_combo.clear()
        for r in self.ink_repo.load_all():
            self.take_ink_combo.addItem(f"[{r['code']}] {r['name']}", r["code"])
        if current:
            idx = self.take_ink_combo.findData(current)
            if idx >= 0:
                self.take_ink_combo.setCurrentIndex(idx)
        self.take_ink_combo.blockSignals(False)
        self._on_stocktake_ink_changed()

    def _on_stocktake_ink_changed(self):
        code = self.take_ink_combo.currentData()
        if not code:
            self.take_system_label.setText("--")
            return
        record = self.ink_repo.get_by_code(code)
        stock = record.get("stock_ml", 0.0) if record else 0.0
        self.take_system_label.setText(f"{stock:g}")
        self.take_actual_spin.setValue(stock)
        self._update_stocktake_diff()

    def _update_stocktake_diff(self):
        try:
            system_val = float(self.take_system_label.text())
        except ValueError:
            system_val = 0.0
        diff = self.take_actual_spin.value() - system_val
        if diff == 0:
            self.take_diff_label.setText("差异：0（账实相符）")
            self.take_diff_label.setStyleSheet("font-weight:700; font-size:14px; color:#16a34a;")
        elif diff > 0:
            self.take_diff_label.setText(f"差异：+{diff:g} ml（实际比账面多）")
            self.take_diff_label.setStyleSheet("font-weight:700; font-size:14px; color:#2563eb;")
        else:
            self.take_diff_label.setText(f"差异：{diff:g} ml（实际比账面少）")
            self.take_diff_label.setStyleSheet("font-weight:700; font-size:14px; color:#dc2626;")

    def _on_confirm_stocktake(self):
        code = self.take_ink_combo.currentData()
        if not code:
            QMessageBox.warning(self, "提示", "请先选择要盘点的油墨。")
            return
        record = self.ink_repo.get_by_code(code)
        system_val = record.get("stock_ml", 0.0) if record else 0.0
        actual_val = self.take_actual_spin.value()
        diff = actual_val - system_val
        if diff == 0:
            QMessageBox.information(self, "提示", "账实相符，不需要调整。")
            return

        self.ink_repo.adjust_stock(code, diff)

        now = datetime.datetime.now()
        self.ink_stock_repo.upsert({
            "id": f"IS-{now.strftime('%Y%m%d%H%M%S%f')}",
            "date": now.strftime("%Y-%m-%d"),
            "movement_type": "入库" if diff > 0 else "出库",
            "ink_code": code,
            "ink_name": record["name"] if record else code,
            "quantity_ml": abs(diff),
            "price": 0,
            "amount": 0,
            "supplier": "",
            "customer": "",
            "notes": f"库存盘点调整（账面{system_val:g} -> 实际{actual_val:g}）",
            "created_at": now.strftime("%Y-%m-%d %H:%M"),
        })
        # 盘点也是改库存，同样要通知【油墨库管理】页刷新（问题3）
        self.ink_db_changed.emit()
        QMessageBox.information(self, "已调整", f"库存已调整为实际数量 {actual_val:g} ml，并记录了一条盘点调整流水。")
        self.refresh_all()

    # ================================================================
    # ③ 客户销售统计
    # ================================================================

    def _build_sales_summary_tab(self) -> QWidget:
        container = QWidget()
        lay = QVBoxLayout(container)
        card = Card("按客户汇总销售情况")
        clay = card.body_layout
        self.sales_table = QTableWidget(0, 4)
        self.sales_table.setHorizontalHeaderLabels(["客户", "累计销售量(ml)", "累计销售金额(元)", "涉及油墨种类数"])
        self.sales_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        clay.addWidget(self.sales_table, 1)
        lay.addWidget(card, 1)
        return container

    def _refresh_sales_summary(self):
        records = [r for r in self.ink_stock_repo.load_all() if r.get("movement_type") == "出库" and r.get("customer")]
        groups: Dict[str, Dict] = {}
        for r in records:
            customer = r.get("customer", "")
            g = groups.setdefault(customer, {"qty": 0.0, "amount": 0.0, "inks": set()})
            g["qty"] += r.get("quantity_ml", 0) or 0
            g["amount"] += r.get("amount", 0) or 0
            g["inks"].add(r.get("ink_code", ""))

        rows = sorted(groups.items(), key=lambda kv: -kv[1]["amount"])
        self.sales_table.setRowCount(len(rows))
        for row, (customer, g) in enumerate(rows):
            items = [
                QTableWidgetItem(customer),
                QTableWidgetItem(f"{g['qty']:g}"),
                QTableWidgetItem(f"{g['amount']:.2f}"),
                QTableWidgetItem(str(len(g["inks"]))),
            ]
            for col, item in enumerate(items):
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.sales_table.setItem(row, col, item)

    # ================================================================
    # 统一刷新 + 顶部统计
    # ================================================================

    def refresh_all(self):
        self._refresh_stock_overview()
        self._refresh_stock_ink_combo()
        self._refresh_stocktake_ink_combo()
        self._refresh_stock_table()
        self._refresh_sales_summary()
        self._refresh_top_stats()

    def _refresh_top_stats(self):
        records = self.ink_stock_repo.load_all()
        purchase_total = sum(r.get("amount", 0) or 0 for r in records if r.get("movement_type") == "入库")
        sales_total = sum(r.get("amount", 0) or 0 for r in records if r.get("movement_type") == "出库" and r.get("customer"))
        self.stat_purchase.set_value(f"¥{purchase_total:,.0f}")
        self.stat_sales.set_value(f"¥{sales_total:,.0f}")

        ink_records = self.ink_repo.load_all()
        stock_value = 0.0
        low_count = 0
        for r in ink_records:
            stock = r.get("stock_ml", 0.0) or 0.0
            recent_price = next(
                (rec.get("price", 0) for rec in records
                 if rec.get("ink_code") == r["code"] and rec.get("movement_type") == "入库" and rec.get("price")),
                None,
            )
            if recent_price:
                stock_value += stock * recent_price
            if self.ink_repo.stock_status(r) in ("缺货", "库存不足"):
                low_count += 1
        self.stat_stock_value.set_value(f"¥{stock_value:,.0f}")
        self.stat_low_stock.set_value(str(low_count))
