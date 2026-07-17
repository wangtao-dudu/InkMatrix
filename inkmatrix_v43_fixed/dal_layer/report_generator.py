# -*- coding: utf-8 -*-
"""
report_generator.py  (DAL层 - 工单文档导出)
=================================================
把"导出数据"这件事当作一种输出侧的数据持久化操作来看待（跟读写JSON
的性质类似，只是格式换成了.docx），所以放在DAL层，与UI/算法解耦：
本模块只接收纯数据结构，不感知PyQt。

生成的Word工单文档参考同行业专色配方通知单的常见格式：
    - 顶部公司抬头（Logo + 公司名称 + 联系方式）
    - 工单基本信息表（产品名称/客户/瓶型/日期）
    - 每个色位一段：目标色色块预览 + ΔE + 配方克重表
"""

import datetime
from typing import Dict, List, Optional

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_cell_background(cell, rgb: tuple):
    """给表格单元格设置背景填充色（python-docx没有直接API，需操作底层XML）。"""
    shading = OxmlElement("w:shd")
    hex_color = "%02X%02X%02X" % rgb
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_heading(doc: Document, text: str, size: int = 14, bold: bool = True, color: Optional[tuple] = None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def generate_print_job_docx(
    company_info: Dict,
    job_meta: Dict,
    stations: List[Dict],
    output_path: str,
) -> str:
    """
    生成多色印刷工单Word文档。

    Args:
        company_info: {"company_name","contact_phone","address","slogan","logo_path"}
        job_meta: {"product_name","customer","bottle_name","operator","order_no"}（均可为空字符串）
        stations: [
            {
                "name": "文字",
                "target_lab": (L,a,b),
                "target_rgb": (r,g,b),          # 用于色块预览底色
                "predicted_lab": (L,a,b) | None,
                "delta_e": float | None,
                "total_grams": float,
                "rows": [{"ink_name":..., "pct":..., "grams":...}, ...],
            }, ...
        ]
        output_path: 输出的.docx文件完整路径

    Returns:
        实际写出的文件路径（等于 output_path）
    """
    doc = Document()

    # ---------------- 公司抬头 ----------------
    logo_path = (company_info or {}).get("logo_path") or ""
    if logo_path:
        try:
            doc.add_picture(logo_path, height=Cm(1.6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
        except Exception:
            pass  # Logo文件读取失败时静默跳过，不影响文档其余内容生成

    company_name = (company_info or {}).get("company_name") or "（未设置公司名称，请在【公司信息】页面填写）"
    _add_heading(doc, company_name, size=18, bold=True)

    contact_line_parts = []
    if (company_info or {}).get("address"):
        contact_line_parts.append(company_info["address"])
    if (company_info or {}).get("contact_phone"):
        contact_line_parts.append(f"电话：{company_info['contact_phone']}")
    if contact_line_parts:
        p = doc.add_paragraph(" · ".join(contact_line_parts))
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    if (company_info or {}).get("slogan"):
        p = doc.add_paragraph(company_info["slogan"])
        p.runs[0].font.size = Pt(9)
        p.runs[0].italic = True
        p.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # 分隔线（用段落下边框模拟，不用表格模拟横线）
    border_p = doc.add_paragraph()
    pPr = border_p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "CBD5E1")
    pBdr.append(bottom)
    pPr.append(pBdr)

    # ---------------- 工单标题与元信息 ----------------
    _add_heading(doc, "多色印刷配方工单 / Print Job Formula Sheet", size=15, bold=True, color=(0x0D, 0x94, 0x88))

    meta = job_meta or {}
    today = datetime.date.today().strftime("%Y-%m-%d")
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_rows = [
        ("产品名称", meta.get("product_name") or "-"),
        ("客户", meta.get("customer") or "-"),
        ("瓶型 / 承印物", meta.get("bottle_name") or "-"),
        ("工单号", meta.get("order_no") or "-"),
        ("制单日期", today),
        ("操作员", meta.get("operator") or "-"),
    ]
    for label, value in meta_rows:
        row = meta_table.add_row()
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[1].text = str(value)

    doc.add_paragraph("")

    # ---------------- 各色位配方 ----------------
    for idx, station in enumerate(stations, start=1):
        _add_heading(doc, f"色位 {idx}：{station.get('name', '')}", size=13, bold=True)

        info_table = doc.add_table(rows=1, cols=3)
        info_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        hdr = info_table.rows[0].cells
        hdr[0].text = "目标色预览"
        hdr[1].text = "目标 Lab"
        hdr[2].text = "预测ΔE"

        row = info_table.add_row()
        rgb = station.get("target_rgb", (255, 255, 255))
        _set_cell_background(row.cells[0], rgb)
        row.cells[0].text = " "
        L, a, b = station.get("target_lab", (0, 0, 0))
        row.cells[1].text = f"({L:.2f}, {a:.2f}, {b:.2f})"
        de = station.get("delta_e")
        de_text = f"{de:.2f}" if de is not None else "未寻优"
        row.cells[2].text = de_text
        if de is not None and de > 3.0:
            row.cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
            row.cells[2].paragraphs[0].runs[0].font.bold = True

        doc.add_paragraph(f"配墨总重：{station.get('total_grams', 0):.2f} g")

        rows_data = station.get("rows", [])
        if rows_data:
            formula_table = doc.add_table(rows=1, cols=3)
            formula_table.style = "Light Grid Accent 1"
            hdr2 = formula_table.rows[0].cells
            hdr2[0].text = "油墨名称"
            hdr2[1].text = "配方比例(%)"
            hdr2[2].text = "精准称重(g)"
            for r in rows_data:
                row2 = formula_table.add_row().cells
                row2[0].text = str(r.get("ink_name", ""))
                row2[1].text = f"{r.get('pct', 0):.2f}%"
                row2[2].text = f"{r.get('grams', 0):.2f} g"
        else:
            doc.add_paragraph("（该色位尚未完成寻优，暂无配方数据）")

        doc.add_paragraph("")

    footer = doc.add_paragraph(f"本文档由 墨算 InkMatrix 智能调色系统 自动生成 · {today}")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)

    doc.save(output_path)
    return output_path
